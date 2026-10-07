import json
import logging
import os
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch


class DependencyCompatibilityTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        environment = patch.dict(os.environ, {"OPENAI_API_KEY": "offline-test-key"}, clear=True)
        environment.start()
        self.addCleanup(environment.stop)
        for target in ("socket.create_connection", "socket.socket.connect"):
            network = patch(target, side_effect=AssertionError("Network access is forbidden in this test"))
            network.start()
            self.addCleanup(network.stop)
        previous_level = logging.root.manager.disable
        logging.disable(logging.CRITICAL)
        self.addCleanup(logging.disable, previous_level)

    def test_h11_rejects_invalid_chunk_terminator(self):
        import h11

        connection = h11.Connection(h11.CLIENT)
        connection.send(h11.Request(method="GET", target="/", headers=[("Host", "example.test")]))
        connection.send(h11.EndOfMessage())
        connection.receive_data(
            b"HTTP/1.1 200 OK\r\nTransfer-Encoding: chunked\r\n\r\n5\r\nhelloXX0\r\n\r\n"
        )
        self.assertIsInstance(connection.next_event(), h11.Response)
        with self.assertRaises(h11.RemoteProtocolError):
            while connection.next_event() is not h11.NEED_DATA:
                pass

    def test_word_loading_and_document_search(self):
        from docx import Document
        from classes.question_answer_base import QuestionAnswerBase
        from handlers.file_handler import load_text_from_word

        paragraph = "POSTHUMAN validators secure the blockchain and earn rewards for their reliable work."
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "documentation.docx"
            document = Document()
            document.add_paragraph(paragraph)
            document.add_paragraph("Delegators choose a validator and can delegate their tokens to support the network.")
            document.save(path)
            self.assertIn(paragraph, load_text_from_word(str(path)))
            context = SimpleNamespace(get_context=lambda: [{"role": "user", "content": paragraph}])
            qa = QuestionAnswerBase(context, str(path))
            self.assertIsNotNone(qa.documentation_model)
            self.assertEqual(qa.search_in_word_document(paragraph), paragraph)

    def test_legacy_openai_request_and_response(self):
        import requests
        from handlers.api_handler import send_to_openai

        response = requests.Response()
        response.status_code = 200
        response.headers["Content-Type"] = "application/json"
        response._content = json.dumps({
            "id": "offline-completion",
            "object": "chat.completion",
            "choices": [{"index": 0, "message": {"role": "assistant", "content": " Offline answer "}, "finish_reason": "stop"}],
        }).encode()
        messages = [{"role": "user", "content": "Offline test"}]
        with patch.object(requests.Session, "request", return_value=response) as request:
            self.assertEqual(send_to_openai(messages), "Offline answer")
        request.assert_called_once()
        payload = json.loads(request.call_args.kwargs["data"])
        self.assertEqual(payload["messages"], messages)
        self.assertEqual(payload["max_tokens"], 750)

    async def test_telegram_initialization_and_handler_registration(self):
        from classes.chat_bot import ChatBot
        from telegram.request import HTTPXRequest

        bot = ChatBot("123456789:offline-test-token")
        self.assertTrue(bot.documentation_text)
        with patch.object(bot.application, "run_polling") as polling:
            bot.start()
        polling.assert_called_once_with()
        self.assertEqual(len(bot.application.handlers[0]), 3)
        payload = json.dumps({"ok": True, "result": {"id": 123456789, "is_bot": True, "first_name": "Offline", "username": "OfflineBot"}}).encode()
        with patch.object(HTTPXRequest, "do_request", new=AsyncMock(return_value=(200, payload))) as request:
            await bot.application.initialize()
            self.assertEqual(bot.application.bot.username, "OfflineBot")
            await bot.application.shutdown()
        self.assertTrue(request.called)

    async def test_message_handler_preserves_context_and_replies(self):
        from classes.chat_bot import ChatBot
        from classes.user_context import UserContext

        bot = ChatBot("123456789:offline-test-token")
        message = SimpleNamespace(
            chat=SimpleNamespace(username="offline_user", first_name="Offline"),
            chat_id=123,
            text="What is POSTHUMAN?",
            reply_text=AsyncMock(),
        )
        context = SimpleNamespace(user_data={"context_memory": [{"role": "system", "content": "Offline documentation"}]})
        with patch("classes.question_answer_base.send_to_openai", return_value="Offline answer") as completion:
            with patch.object(UserContext, "save_context") as save:
                await bot.handle_message(SimpleNamespace(message=message), context)
        message.reply_text.assert_awaited_once_with("Offline answer")
        completion.assert_called_once()
        save.assert_called_once_with(n_messages=2)
        self.assertEqual(context.user_data["context_memory"][-1], {"role": "assistant", "content": "Offline answer"})


if __name__ == "__main__":
    unittest.main()
